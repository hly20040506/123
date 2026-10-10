import json, re, sys
from collections import Counter
import openpyxl

W = '/home/user/123/work'
wb = openpyxl.load_workbook(sys.argv[1])
L = json.load(open(f'{W}/list.json'))
ws = wb.active
ok = True
def chk(cond, msg):
    global ok
    print(('OK   ' if cond else 'FAIL ') + msg)
    ok &= bool(cond)

chk(wb.sheetnames == ['论文汇总'], f'single sheet {wb.sheetnames}')
chk(ws.max_row == 201, f'rows={ws.max_row}')
chk(ws.max_column == 7, f'cols={ws.max_column}')
chk(ws.freeze_panes == 'A2', f'freeze={ws.freeze_panes}')
titles = [ws.cell(r, 1).value for r in range(2, 202)]
chk(len(set(titles)) == 200, f'unique titles={len(set(titles))}')
chk(set(titles) == {x['title'] for x in L}, 'titles match list exactly (no add/drop)')
pdfs = {ws.cell(r, 1).hyperlink.target for r in range(2, 202) if ws.cell(r, 1).hyperlink}
chk(pdfs == {x['pdf'] for x in L}, f'title links = official PDFs ({len(pdfs)})')
for r in range(2, 202):
    t = ws.cell(r, 1).value
    x = next(x for x in L if x['title'] == t)
    if ws.cell(r, 1).hyperlink.target != x['pdf']:
        chk(False, f'row {r} title link mismatch')
fills = colors = 0
for row in ws.iter_rows(min_row=1, max_row=201, max_col=7):
    for c in row:
        if c.fill and c.fill.fill_type:
            fills += 1
        col = c.font.color.rgb if c.font and c.font.color else None
        if col not in (None, 'FF000000'):
            colors += 1
        if not c.alignment.wrap_text:
            chk(False, f'no wrap {c.coordinate}')
chk(fills == 0, f'no fills ({fills})')
chk(colors == 0, f'all black font ({colors})')
C = Counter(ws.cell(r, 3).value for r in range(2, 202))
chk(set(C) <= {'有', '无'}, f'proc values {dict(C)}')
D = Counter(ws.cell(r, 4).value for r in range(2, 202))
chk(set(D) <= {'是（公开代码）', '仅仓库页面，未确认代码', '未找到公开代码', '代码链接暂无法确认'}, f'open-source values {dict(D)}')
for r in range(2, 202):
    d, f = ws.cell(r, 4), ws.cell(r, 6)
    if d.value != '未找到公开代码':
        if not (d.hyperlink and f.hyperlink and d.hyperlink.target == f.hyperlink.target): chk(False, f"row {r} code link missing")
    pair = {'是（公开代码）': '是（提供代码）', '仅仓库页面，未确认代码': '未确认代码', '未找到公开代码': '未找到代码', '代码链接暂无法确认': '代码链接待确认'}[d.value]
    chk(f.value.startswith(pair), f'row {r} D/F consistent') if not f.value.startswith(pair) else None
timing = re.compile(r'(\d\s*ms\b|毫秒|FPS(?!）)|[Ll]atency|耗时|推理时间|训练时间|运行时间|[Rr]untime|s/frame|min/|GPU\s?hours?|AITS)')
bad = []
for r in range(2, 202):
    for c in (5, 7):
        v = str(ws.cell(r, c).value)
        for m in timing.finditer(v):
            ctx = v[max(0, m.start() - 25):m.end() + 10].replace('\n', ' / ')
            bad.append(f'{ws.cell(r,1).value[:30]} | {ctx}')
print('timing-like hits to review:', len(bad))
for b in bad:
    print('   ', b)
empty = [r for r in range(2, 202) for c in range(1, 8) if not ws.cell(r, c).value]
chk(not empty, f'no empty cells ({len(empty)})')
nonum = [r for r in range(2, 202) if not re.search(r'\d', re.sub(r'（表\d+.*?）|p\.\d+', '', ws.cell(r, 7).value))]
print('metrics cells without any number:', len(nonum))
print('ALL OK' if ok else 'SOME CHECKS FAILED')
