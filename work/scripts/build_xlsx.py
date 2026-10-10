import json, re, sys, math, unicodedata
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.hyperlink import Hyperlink

W = '/home/user/123/work'
OUT = sys.argv[1]
L = json.load(open(f'{W}/list.json'))
R = [json.load(open(f'{W}/results/{i:03d}.json')) for i in range(1, 201)]

DOM_MAP = {'视觉语言导航': '具身导航', '自动驾驶数据集与规划': '自动驾驶规划'}

def dom(d):
    a, b = d['dom'].split('\n', 1)
    return DOM_MAP.get(a, a), b

def main_ds(m):
    for line in m.split('\n'):
        if line.startswith(('数据集/基准论文', '注', '数据处理')):
            continue
        return re.split(r'[（(：:，, ]', line, maxsplit=1)[0].strip()
    return ''

def code_cells(d):
    u = d.get('url')
    st = d.get('repo_state')
    if u and st == 'code':
        f = '是（提供代码）'
        if d.get('repo_note'):
            f += '\n注：' + d['repo_note']
        return '是（公开代码）', f, u
    if u and st == 'readme':
        return '仅仓库页面，未确认代码', '未确认代码（仓库仅有README/项目页文件）', u
    if u and st == 'inaccessible':
        return '代码链接暂无法确认', '代码链接待确认（链接无法访问）', u
    if d.get('soon'):
        return '未找到公开代码', '未找到代码（' + d.get('soon_text', '作者/项目页注明代码将发布') + '）', None
    return '未找到公开代码', '未找到代码', None

rows = []
for i, (x, d) in enumerate(zip(L, R), 1):
    assert d['idx'] == i
    a, b = dom(d)
    rows.append(dict(i=i, x=x, d=d, major=a, minor=b, ds=main_ds(d['metrics'])))

cnt = {}
for r in rows:
    cnt[r['major']] = cnt.get(r['major'], 0) + 1
rows.sort(key=lambda r: (-cnt[r['major']], r['major'], r['ds'].lower(), r['i']))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = '论文汇总'
HDR = ['论文题目', '所属领域', '数据特殊处理\n（有/无）', '是否开源（GitHub）', '显卡型号', '能否复现\n（按提供代码）', '数据集与核心指标（作者方法）']
fh = Font(name='Microsoft YaHei', size=11, bold=True, color='FF000000')
fd = Font(name='Microsoft YaHei', size=11, color='FF000000')
ah = Alignment(horizontal='center', vertical='center', wrap_text=True)
ad = Alignment(horizontal='left', vertical='center', wrap_text=True)
for c, h in enumerate(HDR, 1):
    cell = ws.cell(1, c, h)
    cell.font = fh
    cell.alignment = ah
WID = {'A': 57, 'B': 25, 'C': 18, 'D': 27, 'E': 25, 'F': 22, 'G': 150}
for k, v in WID.items():
    ws.column_dimensions[k].width = v
ws.row_dimensions[1].height = 52
ws.freeze_panes = 'A2'

def dispw(s):
    return sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in s)

def nlines(s, width):
    return sum(max(1, math.ceil(dispw(line) / (width * 0.9))) for line in str(s).split("\n"))

for r_i, r in enumerate(rows, 2):
    x, d = r['x'], r['d']
    pdf = x['pdf']
    D, F, u = code_cells(d)
    vals = [x['title'], r['major'] + '\n' + r['minor'], '无' if d['proc'] == '未说明' else d['proc'], D, d['gpu'], F, d['metrics']]
    links = [pdf, None,
             f"{pdf}#page={d['proc_p']}" if d.get('proc_p') else None,
             u, None, u, f"{pdf}#page={d['mp']}" if d.get('mp') else pdf]
    for c, (v, lk) in enumerate(zip(vals, links), 1):
        cell = ws.cell(r_i, c, v)
        cell.font = fd
        cell.alignment = ad
        if lk:
            cell.hyperlink = lk
            cell.font = fd  # keep black, no underline
    h = max(nlines(v, WID[k]) for v, k in zip(vals, 'ABCDEFG'))
    ws.row_dimensions[r_i].height = max(64, min(409, 18 * h + 14))

wb.save(OUT)
print('saved', OUT, ws.max_row, ws.max_column)
